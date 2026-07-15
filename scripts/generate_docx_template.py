import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_resume_template():
    doc = docx.Document()
    
    # Page setup - 0.5 inch margins for maximum printable area
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Base style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(9.0)  # Standard size for detailed single-page resumes

    # 1. Header (Name, Specialization)
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    
    # Widths: Left (for Photo) ~ 1.2 in, Right (for Name/Course) ~ 6.3 in
    header_table.columns[0].width = Inches(1.2)
    header_table.columns[1].width = Inches(6.3)
    
    # Prevent table borders
    tblPr = header_table._tbl.tblPr
    borders = parse_xml(
        '<w:tblBorders %s>'
        '<w:top w:val="none"/>'
        '<w:left w:val="none"/>'
        '<w:bottom w:val="none"/>'
        '<w:right w:val="none"/>'
        '<w:insideH w:val="none"/>'
        '<w:insideV w:val="none"/>'
        '</w:tblBorders>' % nsdecls('w')
    )
    tblPr.append(borders)

    # Left cell: Photo Placeholder (Compact)
    cell_left = header_table.cell(0, 0)
    cell_left.width = Inches(1.2)
    p_photo = cell_left.paragraphs[0]
    p_photo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_photo.paragraph_format.space_before = Pt(4)
    p_photo.paragraph_format.space_after = Pt(4)
    run_photo = p_photo.add_run("[ INSERT\nPHOTO\nHERE ]")
    run_photo.font.bold = True
    run_photo.font.size = Pt(8)
    run_photo.font.color.rgb = RGBColor(120, 120, 120)
    
    tcPr = cell_left._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        '<w:tcBorders %s>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '</w:tcBorders>' % nsdecls('w')
    )
    tcPr.append(tcBorders)

    # Right cell: Name and Title (Compact)
    cell_right = header_table.cell(0, 1)
    cell_right.width = Inches(6.3)
    p_name = cell_right.paragraphs[0]
    p_name.paragraph_format.space_before = Pt(4)
    p_name.paragraph_format.space_after = Pt(2)
    p_name.paragraph_format.left_indent = Inches(0.15)
    run_name = p_name.add_run("[YOUR FULL NAME]")
    run_name.font.size = Pt(16)  # Slightly smaller to avoid wrap
    run_name.font.bold = True
    run_name.font.color.rgb = RGBColor(11, 86, 148)

    p_degree = cell_right.add_paragraph()
    p_degree.paragraph_format.space_before = Pt(0)
    p_degree.paragraph_format.space_after = Pt(0)
    p_degree.paragraph_format.left_indent = Inches(0.15)
    run_degree = p_degree.add_run("[YOUR DEGREE / PROGRAM, e.g., Bachelor of Science in Information Technology]")
    run_degree.font.size = Pt(10)
    run_degree.font.color.rgb = RGBColor(80, 80, 80)

    # Add a horizontal line below header (using paragraph bottom border with minimal height)
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(4)
    p_space.paragraph_format.space_after = Pt(8)
    pPr = p_space._p.get_or_add_pPr()
    pBdr = parse_xml(
        '<w:pBdr %s>'
        '<w:bottom w:val="single" w:sz="6" w:space="1" w:color="0B5694"/>'
        '</w:pBdr>' % nsdecls('w')
    )
    pPr.append(pBdr)

    # 2. Main Two-Column Layout Table
    main_table = doc.add_table(rows=1, cols=2)
    main_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    main_table.autofit = False
    
    # Column widths: Left ~ 4.7 in, Right ~ 2.8 in
    main_table.columns[0].width = Inches(4.7)
    main_table.columns[1].width = Inches(2.8)

    # Remove main table borders
    tblPr_main = main_table._tbl.tblPr
    borders_main = parse_xml(
        '<w:tblBorders %s>'
        '<w:top w:val="none"/>'
        '<w:left w:val="none"/>'
        '<w:bottom w:val="none"/>'
        '<w:right w:val="none"/>'
        '<w:insideH w:val="none"/>'
        '<w:insideV w:val="none"/>'
        '</w:tblBorders>' % nsdecls('w')
    )
    tblPr_main.append(borders_main)

    col_left = main_table.cell(0, 0)
    col_left.width = Inches(4.7)
    col_right = main_table.cell(0, 1)
    col_right.width = Inches(2.8)

    # Helper function to add structured headers (Compact)
    def add_section_header(cell, text):
        p = cell.paragraphs[0] if len(cell.paragraphs[0].text) == 0 else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(11.0)
        run.font.bold = True
        run.font.color.rgb = RGBColor(11, 86, 148)
        
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(
            '<w:pBdr %s>'
            '<w:bottom w:val="single" w:sz="4" w:space="1" w:color="CCCCCC"/>'
            '</w:pBdr>' % nsdecls('w')
        )
        pPr.append(pBdr)
        return p

    # --- LEFT COLUMN CONTENT ---
    # Section: Contact Information
    add_section_header(col_left, "Contact Information")
    
    p = col_left.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run("Phone:\t")
    run.font.bold = True
    p.add_run("[CONTACT NUMBER, e.g., 09303044030]")
    
    p = col_left.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run("Email:\t")
    run.font.bold = True
    p.add_run("[EMAIL ADDRESS, e.g., student@emaildomain.com]")
    
    p = col_left.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("Address:\t")
    run.font.bold = True
    p.add_run("[PHYSICAL ADDRESS, e.g., Quezon City, Metro Manila 1109]")

    # Section: Education
    add_section_header(col_left, "Education")
    
    p = col_left.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run("[DEGREE / PROGRAM DETAILS, e.g., BS in Information Technology]")
    run.font.bold = True
    run.font.size = Pt(9.5)
    
    p = col_left.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run("[YOUR UNIVERSITY / ACADEMIC INSTITUTION, e.g., FEU Institute of Technology]")
    run.font.bold = True
    run.font.color.rgb = RGBColor(80, 80, 80)
    
    p = col_left.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.add_run("[CAMPUS LOCATION, e.g., P. Paredes St., Sampaloc, Manila]")
    
    p = col_left.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("Expected Graduation – [Year, e.g., 2025]")
    run.font.italic = True

    # Section: Work Experience (Optional)
    add_section_header(col_left, "Work Experience (Optional)")
    
    p = col_left.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run("[Job Title / Position, e.g., Web Developer Intern]")
    run.font.bold = True
    run.font.size = Pt(9.5)
    
    p = col_left.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run("[Company / Organization Name, e.g., Tech Solutions Inc.]")
    run.font.bold = True
    run.font.color.rgb = RGBColor(80, 80, 80)
    
    p = col_left.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run("[Inclusive Dates, e.g., June 2024 – August 2024] | ")
    run.font.italic = True
    p.add_run("[Location, e.g., Makati City]")
    
    p = col_left.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(1)
    p.add_run("[Describe responsibility/achievement, e.g., Developed responsive UI components using Alpine.js]")
    
    p = col_left.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(8)
    p.add_run("[Describe another achievement, e.g., Integrated RESTful API endpoints with Laravel backend]")

    # Section: Key Skills / Proficiency
    add_section_header(col_left, "Key Skills / Proficiency")
    
    p = col_left.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("Technical Skills:")
    run.font.bold = True
    
    p = col_left.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run("Programming Languages | ")
    run.font.bold = True
    p.add_run("HTML, JavaScript, Kotlin, PHP, Python")
    
    p = col_left.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run("Database & Dev Tools | ")
    run.font.bold = True
    p.add_run("SQL, Git, Azure, AWS")
    
    p = col_left.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    p.add_run("UI/UX Design & Prototyping (Figma)")
    
    p = col_left.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("Soft Skills:")
    run.font.bold = True
    
    p = col_left.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(1)
    p.add_run("Oral and written communication")
    p = col_left.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(8)
    p.add_run("Teamwork & Collaboration")

    # Section: Accomplished Projects / Engagements (Compact - 2 items)
    add_section_header(col_left, "Accomplished Projects / Engagements")
    
    p = col_left.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    run_title = p.add_run("[Project Title 1]")
    run_title.font.bold = True
    p.add_run(" : [Description of web/mobile system project] | ")
    run_role = p.add_run("[Role, e.g., Lead Developer]")
    run_role.font.italic = True
    
    p = col_left.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(0)
    run_title = p.add_run("[Project Title 2]")
    run_title.font.bold = True
    p.add_run(" : [Description of programming course project] | ")
    run_role = p.add_run("[Role, e.g., Project Manager]")
    run_role.font.italic = True


    # --- RIGHT COLUMN CONTENT ---
    # Section: Awards
    add_section_header(col_right, "Awards")
    
    p = col_right.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(1)
    p.add_run("[Award Title 1 / Competition Name]")
    
    p = col_right.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(8)
    p.add_run("[Award Title 2 / Competition Name]")

    # Section: Certifications
    add_section_header(col_right, "Certifications")
    
    p = col_right.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(1)
    p.add_run("[Certification Title 1, e.g., Microsoft Certified]")
    
    p = col_right.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(8)
    p.add_run("[Certification Title 2]")

    # Section: Organizations
    add_section_header(col_right, "Organizations")
    
    p = col_right.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run("[Officer Position]")
    run.font.bold = True
    p.add_run(", [Organization Name]")
    
    p = col_right.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("Member")
    run.font.bold = True
    p.add_run(", [Organization Name]")

    # Section: Languages Spoken
    add_section_header(col_right, "Languages Spoken")
    
    p = col_right.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(1)
    p.add_run("Filipino")
    p = col_right.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(8)
    p.add_run("English")

    # Section: Other Interests
    add_section_header(col_right, "Other Interests")
    
    p = col_right.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(1)
    p.add_run("[Hobby 1, e.g., Vlogging]")
    p = col_right.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(0)
    p.add_run("[Hobby 2, e.g., Digital Painting]")

    doc.save("public/templates/SmarTern_Student_Resume_Template.docx")
    print("Template generated successfully!")

if __name__ == '__main__':
    create_resume_template()
